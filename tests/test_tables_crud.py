"""
Tables CRUD API testlari.
"""

from io import BytesIO

import pytest
from fastapi.testclient import TestClient

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill


@pytest.mark.api
class TestTablesCRUD:
    """Tables CRUD testlari."""

    def _create_chat(self, client: TestClient) -> int:
        response = client.post("/api/v1/chats", json={"title": "Test"})
        return response.json()["id"]

    def test_upload_csv_table(self, client: TestClient):
        chat_id = client.post(
            "/api/v1/chats",
            json={"title": "Yangi chat"},
        ).json()["id"]

        response = client.post(
            f"/api/v1/chats/{chat_id}/table/upload",
            files={
                "file": (
                    "xodimlar.csv",
                    "Ism,Maosh\nAli,5000\nVali,4200\n".encode(),
                    "text/csv",
                ),
            },
        )

        assert response.status_code == 200
        table = response.json()
        assert table["title"] == "xodimlar"
        assert table["columns"] == ["Ism", "Maosh"]
        assert table["rows"] == [["Ali", "5000"], ["Vali", "4200"]]
        assert table["source"] == "upload"
        assert table["styles"]["body"]["border_color"] == "#000000"

        assert client.get(f"/api/v1/chats/{chat_id}").json()["title"] == "xodimlar"
        messages = client.get(f"/api/v1/chats/{chat_id}/messages").json()["messages"]
        assert [message["role"] for message in messages] == ["user"]

    def test_upload_xlsx_table(self, client: TestClient):
        chat_id = self._create_chat(client)
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Mahsulotlar"
        worksheet.append(["Nomi", "Narx"])
        worksheet.append(["Daftar", 12000])
        worksheet.append(["Ruchka", 5000])
        worksheet["A1"].fill = PatternFill(
            fill_type="solid",
            fgColor="FF123456",
        )
        worksheet["B2"].fill = PatternFill(
            fill_type="solid",
            fgColor="FF16A34A",
        )
        content = BytesIO()
        workbook.save(content)

        response = client.post(
            f"/api/v1/chats/{chat_id}/table/upload",
            files={
                "file": (
                    "mahsulotlar.xlsx",
                    content.getvalue(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            },
        )

        assert response.status_code == 200
        assert response.json()["title"] == "Mahsulotlar"
        assert response.json()["columns"] == ["Nomi", "Narx"]
        assert response.json()["rows"] == [["Daftar", 12000], ["Ruchka", 5000]]
        assert response.json()["styles"]["cells"]["0:0"]["bg_color"] == "#123456"
        assert response.json()["styles"]["cells"]["1:1"]["bg_color"] == "#16A34A"

        imported = response.json()
        exported = client.post(
            "/api/v1/excel/export",
            json={
                "title": imported["title"],
                "columns": imported["columns"],
                "rows": imported["rows"],
                "styles": imported["styles"],
            },
        )
        assert exported.status_code == 200
        exported_workbook = load_workbook(BytesIO(exported.content))
        exported_sheet = exported_workbook.active
        assert exported_sheet["A1"].fill.fgColor.rgb.endswith("123456")
        assert exported_sheet["B2"].fill.fgColor.rgb.endswith("16A34A")

    def test_upload_rejects_unsupported_file(self, client: TestClient):
        chat_id = self._create_chat(client)

        response = client.post(
            f"/api/v1/chats/{chat_id}/table/upload",
            files={"file": ("table.txt", b"not a table", "text/plain")},
        )

        assert response.status_code == 422
        assert "Faqat .xlsx yoki .csv" in response.json()["detail"]

    def test_create_table(self, client: TestClient):
        """Jadval yaratish ishlashi kerak."""
        chat_id = self._create_chat(client)

        response = client.post(
            f"/api/v1/chats/{chat_id}/table",
            json={
                "title": "Test",
                "columns": ["A", "B"],
                "rows": [["1", "2"]],
            },
        )
        assert response.status_code == 201
        data = response.json()
        assert data["title"] == "Test"
        assert data["row_count"] == 1
        assert data["column_count"] == 2
        assert data["version"] == 1
        assert data["styles"]["body"]["border_color"] == "#000000"
        assert data["styles"]["body"]["border_sides"] == [
            "top",
            "bottom",
            "left",
            "right",
        ]

    def test_styles_are_saved_and_updated(self, client: TestClient):
        chat_id = self._create_chat(client)
        response = client.post(
            f"/api/v1/chats/{chat_id}/table",
            json={
                "title": "Test",
                "columns": ["Maosh"],
                "rows": [[5000]],
                "styles": {"columns": {"Maosh": {"bg_color": "#16A34A"}}},
            },
        )
        assert response.status_code == 201
        assert response.json()["styles"]["columns"]["Maosh"]["bg_color"] == "#16A34A"

        response = client.patch(
            f"/api/v1/chats/{chat_id}/table",
            json={"styles": {"header": {"bg_color": "#112233"}}},
        )
        assert response.status_code == 200
        assert response.json()["styles"]["header"]["bg_color"] == "#112233"
        assert response.json()["styles"]["columns"]["Maosh"]["bg_color"] == "#16A34A"

    def test_get_table(self, client: TestClient):
        """Jadvalni olish."""
        chat_id = self._create_chat(client)
        client.post(
            f"/api/v1/chats/{chat_id}/table",
            json={
                "title": "Test",
                "columns": ["A"],
                "rows": [["1"]],
            },
        )

        response = client.get(f"/api/v1/chats/{chat_id}/table")
        assert response.status_code == 200
        assert response.json()["title"] == "Test"

    def test_create_table_twice_fails(self, client: TestClient):
        """Ikki marta jadval yaratish xato berishi kerak."""
        chat_id = self._create_chat(client)

        client.post(
            f"/api/v1/chats/{chat_id}/table",
            json={"title": "T", "columns": ["A"], "rows": [["1"]]},
        )

        response = client.post(
            f"/api/v1/chats/{chat_id}/table",
            json={"title": "T2", "columns": ["B"], "rows": [["2"]]},
        )
        assert response.status_code == 400

    def test_upsert_table(self, client: TestClient):
        """PUT — yaratish yoki yangilash."""
        chat_id = self._create_chat(client)

        # 1-marta — yaratadi
        response = client.put(
            f"/api/v1/chats/{chat_id}/table",
            json={"title": "T1", "columns": ["A"], "rows": [["1"]]},
        )
        assert response.status_code == 200
        assert response.json()["version"] == 1

        # 2-marta — yangilaydi
        response = client.put(
            f"/api/v1/chats/{chat_id}/table",
            json={"title": "T2", "columns": ["A", "B"], "rows": [["1", "2"]]},
        )
        assert response.status_code == 200
        assert response.json()["title"] == "T2"
        assert response.json()["version"] == 2

    def test_update_table(self, client: TestClient):
        """Jadvalni tahrirlash."""
        chat_id = self._create_chat(client)
        client.post(
            f"/api/v1/chats/{chat_id}/table",
            json={"title": "T", "columns": ["A"], "rows": [["1"]]},
        )

        response = client.patch(
            f"/api/v1/chats/{chat_id}/table",
            json={"rows": [["1"], ["2"]]},
        )
        assert response.status_code == 200
        assert response.json()["row_count"] == 2
        assert response.json()["version"] == 2

    def test_delete_table(self, client: TestClient):
        """Jadvalni o'chirish."""
        chat_id = self._create_chat(client)
        client.post(
            f"/api/v1/chats/{chat_id}/table",
            json={"title": "T", "columns": ["A"], "rows": [["1"]]},
        )

        response = client.delete(f"/api/v1/chats/{chat_id}/table")
        assert response.status_code == 204

        # Tekshirish
        response = client.get(f"/api/v1/chats/{chat_id}/table")
        assert response.status_code == 404

    def test_get_nonexistent_table(self, client: TestClient):
        """Yo'q jadval 404 qaytarishi kerak."""
        chat_id = self._create_chat(client)
        response = client.get(f"/api/v1/chats/{chat_id}/table")
        assert response.status_code == 404

    def test_table_with_formulas(self, client: TestClient):
        """Formulali jadval yaratish."""
        chat_id = self._create_chat(client)

        response = client.post(
            f"/api/v1/chats/{chat_id}/table",
            json={
                "title": "Talabalar",
                "columns": ["Ism", "Ball", "Natija"],
                "rows": [
                    ["Ali", 92, '=IF(B2>60,"O\'tdi","Yiqildi")'],
                    ["Vali", 45, '=IF(B3>60,"O\'tdi","Yiqildi")'],
                ],
            },
        )
        assert response.status_code == 201
        data = response.json()
        assert data["has_formulas"] is True
        assert data["formula_count"] == 2